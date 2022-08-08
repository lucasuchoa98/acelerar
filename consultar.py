from dataclasses import dataclass
from urllib import request
import pandas as pd
import requests
from bs4 import BeautifulSoup as bs
from docx import Document
from datetime import date

today = date.today()

base = r'https://sei.al.gov.br/sei/modulos/pesquisa/md_pesq_processo_exibir.php?'


url = [base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ5mZpq69-u7jcwJJGsy7PsE5roDpEp5uxy4oCg82u659',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ004SFEypb6o2A8QG5-qsafKqdNsMd5NkhgQj5Vci5gT',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ4YP2NUKvA7w03DU9LmXUc7kzQOIT6dR-X6-3h9f1r2m',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ5KP6mECNml0a4XKorQJItQOg2MZ2vFIqwBLTvzLz6tl',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ0-cBeUHrkZFa1uQvo-6IDAEIv9uBSQEuh9Oq8o05v6A',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ5yq79TlrLNq_BL5i_mk89ZlxjbIwGCenOCzelbfSNJE',
    base + r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ0AavRk-GYIH50X_riJ-jz1E2chvBH2clJJdzLakIk2f',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ0PUurfloNfhg_kqoW3ur_bhHXUVosuKqlzi8X58RHjp',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ-g6DLvqo4N6z2ZyVegHRMFuWF3rZmmwDzB8kFSCEGmY',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ-M0_cOJo5B0bYc_1FjThXVHoJUy_MIleaz2CGeD8078',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ_24bL4Snh9Koaw_maGROEhxN7tgmUVL3qbSXOAtL_2z',
    base +r'IC2o8Z7ACQH4LdQ4jJLJzjPBiLtP6l2FsQacllhUf-duzEubalut9yvd8-CzYYNLu7pd-wiM0k633-D6khhQNTrK0DGat35bVWLoXU-NZPMGUY_d_q5nbolcyZtqfgkX',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ5mZpq69-u7jcwJJGsy7PsE5roDpEp5uxy4oCg82u659',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ_r29b6OUyZz4tTn9eynw5Yb46WI-aJFl5_O3AuynVfp',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQxs7ucJ_y8hDBhBOrs_Fa5zdxSgLlAkw8IgW_SbPY-P-',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ0UNW7gsWR0Mp0VqJvgoVqVNpqlunhaAzooXo5f_jE-Z',
    base +r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQxDm5WGst-YAl4k1SWsrAr3RgmicSI66l0LMZQQ3y1u8',
    base+ r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ8qGyiZo_AQ5k0i2enoXtf5-JaQCz61LsnS18wEvXURu',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ7dn9ujj1Iok2cIyyJ5UwyBlLhRfl768sixWDg79fNRy',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ3SVcYIcs1DWN6FvxrGBau8Wk70UK4EdWj9T6BY1enL3',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ3W7Lin1ixc8GQMXLOH_mIQp35B4pUCROaiQcDJqlhE_',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ982KMjxWmtaYLa1hSfvaEBUJrTV5IPhyzs6wLmWDrqQ',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ65OhdOEGEA-8vMBX86m-kJYsCwoHW3y1EI0Hwe7wWgS',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ-H3-9j6PPesXffkNrQysSHgm2leyrpnN9OVN4hcjiO6',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ_Y8OvkKLnujpGR4xGQRyCCbOWRXvpmHRMiYOUlWw1cA',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ_8JAfcp3hz9uPiQotLx26Bjru6yFq8baMcd8V89IvSn',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQ8roCZNvdTd0tJtjgcp7hD8U9xguflTL-YYSzdoJxndP',
    base+r'iI3OtHvPArITY997V09rhsSkbDKbaYSycOHqqF2xsM0IaDkkEyJpus7kCPb435VNEAb16AAxmJKUdrsNWVIqQz8yUpQKgDwQVa145TWFr-EJdY51u3BA_BtOoXd0cqhc',
    ]

def observacao():
    rows = soup.find_all('tr',{'class': 'andamentoAberto'})
    columns = rows.pop()
    columns = columns.text.split('\n')
    data = columns[1]
    unidade = columns[2]
    descricao = columns[3]
    resposta = [data, unidade, descricao]
    
    return resposta

def processo():
    rows = soup.find('tr',{'class': 'infraTrClara'}).get_text()
    rows = rows.split(':')

    return rows

def df_to_word(data: dict, report_name:str):
    assert type(data) == dict
    assert '.docx' in report_name
    df = pd.DataFrame(data)
    doc = Document()

    table = doc.add_table(df.shape[0]+1, df.shape[1])

    for j in range(df.shape[-1]):
        table.cell(0,j).text = df.columns[j]

    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            table.cell(i+1,j).text = str(df.values[i,j])

    doc.save(f'./{report_name}')

protocolo_list = []
observacao_list = []

for i in url:
    try:
        response = requests.get(i)

    except:
        print('fail')
        break

    soup = bs(response.content, "html.parser")

    n_proc = processo()
    status = observacao()

    protocolo_list.append(n_proc[2])
    observacao_list.append(f'{status[2]} {status[1]} em {status[0]}')




dicionario = {
    'Protocolo':protocolo_list,
    'Observacoes': observacao_list
}

dia = today.strftime("%d-%m")
df_to_word(dicionario,f'StatusDosProcessosSEMARH{dia}.docx')


#print(f"Protocolo n° {n_proc[2]} - {status[2]} {status[1]} em {status[0]}")